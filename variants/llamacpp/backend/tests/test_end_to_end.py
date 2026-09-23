import time
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from threading import Event

from docx import Document
from fastapi.testclient import TestClient

from app.main import create_app


def test_reopen_export_and_delete(tmp_path, permitted_wav, monkeypatch):
    db, files = tmp_path / "db.sqlite", tmp_path / "files"

    def processor(*_args):
        return {"transcript": [{"id": "t1", "start": 0.0, "end": 1.0,
                                "text": "Фиксируем аудит", "speaker_id": "speaker_0",
                                "overlap": False, "speaker_uncertain": False}],
                "summary": "Обсудили аудит", "actions": [{
                    "text": "Провести аудит", "assignee": None,
                    "deadline_phrase": None, "due_date": None,
                    "source_turn_ids": ["t1"], "confirmation_turn_ids": ["t1"],
                    "needs_review": True,
                }]}

    def convert(args, **_kwargs):
        output_dir = args[args.index("--outdir") + 1]
        from pathlib import Path
        Path(output_dir, "draft.pdf").write_bytes("%PDF demo Проверено Отдел Б".encode())

    monkeypatch.setattr("app.export.subprocess.run", convert)
    with TestClient(create_app(db, files, processor=processor)) as client:
        response = client.post("/meetings", data={"meeting_date": "2026-09-23"},
                               files={"audio": ("demo.wav", permitted_wav, "audio/wav")})
        meeting_id = response.json()["id"]
        client.post(f"/meetings/{meeting_id}/process")
        for _ in range(100):
            if client.get(f"/meetings/{meeting_id}").json()["status"] == "review":
                break
            time.sleep(0.01)
        else:
            raise AssertionError("No review result")
        action = processor()["actions"][0]
        action.update(assignee="Отдел Б", needs_review=True, confirmation_checked=False)
        saved = client.patch(f"/meetings/{meeting_id}", json={"summary": "Проверено",
                          "actions": [action]})
        assert saved.status_code == 200
        checked = saved.json()["actions"][0]
        checked.update(needs_review=False, confirmation_checked=True)
        assert client.patch(f"/meetings/{meeting_id}", json={"actions": [checked]}).status_code == 200

    with TestClient(create_app(db, files, processor=processor)) as client:
        assert client.get(f"/meetings/{meeting_id}").json()["summary"] == "Проверено"
        assert client.post(f"/meetings/{meeting_id}/export").status_code == 200
        docx = client.get(f"/meetings/{meeting_id}/export/docx")
        pdf = client.get(f"/meetings/{meeting_id}/export/pdf")
        assert "Проверено" in " ".join(p.text for p in Document(BytesIO(docx.content)).paragraphs)
        assert "Отдел Б" in " ".join(c.text for row in Document(BytesIO(docx.content)).tables[0].rows for c in row.cells)
        assert pdf.content.startswith(b"%PDF")
        assert client.patch(f"/meetings/{meeting_id}", json={"summary": "Исправлено"}).status_code == 200
        assert client.get(f"/meetings/{meeting_id}/export/docx").status_code == 404
        assert client.get(f"/meetings/{meeting_id}/export/pdf").status_code == 404
        assert client.delete(f"/meetings/{meeting_id}").status_code == 204
        assert client.get(f"/meetings/{meeting_id}").status_code == 404
        assert not list(files.iterdir())


def test_delete_serializes_with_active_export(tmp_path, permitted_wav, monkeypatch):
    entered, release = Event(), Event()

    def blocked_export(meeting, output, **kwargs):
        entered.set()
        assert release.wait(5)
        from pathlib import Path
        output.mkdir(parents=True, exist_ok=True)
        docx, pdf = Path(output, "draft.docx"), Path(output, "draft.pdf")
        docx.write_bytes(b"docx")
        pdf.write_bytes(b"%PDF demo")
        return docx, pdf

    monkeypatch.setattr("app.main.export_draft", blocked_export)
    with TestClient(create_app(tmp_path / "db.sqlite", tmp_path / "files")) as client:
        created = client.post("/meetings", data={"meeting_date": "2026-09-23"},
                              files={"audio": ("demo.wav", permitted_wav, "audio/wav")})
        meeting_id = created.json()["id"]
        from app.store import Store
        Store(tmp_path / "db.sqlite", tmp_path / "files").finish(meeting_id, {
            "transcript": [], "summary": "", "actions": [],
        })
        with ThreadPoolExecutor(max_workers=2) as pool:
            exporting = pool.submit(client.post, f"/meetings/{meeting_id}/export")
            assert entered.wait(5)
            deleting = pool.submit(client.delete, f"/meetings/{meeting_id}")
            release.set()
            assert exporting.result().status_code == 200
            assert deleting.result().status_code == 204
