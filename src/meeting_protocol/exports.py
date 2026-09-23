from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from .audio import SetupError
from .extraction import display_name
from .models import Meeting


def stamp(seconds: float) -> str:
    value = int(seconds)
    return f"{value // 3600:02}:{value // 60 % 60:02}:{value % 60:02}"


def task_rows(meeting):
    for i, task in enumerate(meeting.tasks, 1):
        yield [
            str(i),
            task.action + (" [NEEDS REVIEW]" if task.needs_review else ""),
            display_name(meeting, task.assignee) or "-",
            str(task.deadline_normalized or task.deadline_raw or "-"),
            task.status.value,
        ]


def export_docx(meeting: Meeting) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document()
    doc.sections[0].left_margin = Inches(0.7)
    doc.sections[0].right_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("MEETING PROTOCOL", 0)
    doc.add_paragraph(meeting.title)
    doc.add_paragraph("Date: " + str(meeting.meeting_date or "Not supplied"))
    doc.add_paragraph("Participants: " + ", ".join(p.name for p in meeting.participants))
    doc.add_heading("1. Summary", 1)
    doc.add_paragraph(meeting.summary.text)
    doc.add_heading("2. Key decisions", 1)
    for decision in meeting.decisions:
        doc.add_paragraph(decision.text, style="List Bullet")
    doc.add_heading("3. Tasks", 1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["#", "Task", "Assignee", "Deadline", "Status"]):
        cell.text = text
    for row in task_rows(meeting):
        for cell, text in zip(table.add_row().cells, row):
            cell.text = text
    doc.add_heading("4. Transcript", 1)
    for segment in meeting.transcript:
        doc.add_paragraph(
            f"[{stamp(segment.start)}] {segment.speaker} / {display_name(meeting, segment.speaker)}\n{segment.text}"
        )
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def export_pdf(meeting: Meeting, font_path: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    candidates = [
        font_path,
        "C:/Windows/Fonts/arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    font = next((Path(p) for p in candidates if p and Path(p).is_file()), None)
    if font is None:
        raise SetupError("Set PDF_FONT_PATH to a Unicode TTF font supporting Cyrillic and Kazakh.")
    pdfmetrics.registerFont(TTFont("MeetingUnicode", str(font)))
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = "MeetingUnicode"

    def paragraph(text, style="Normal"):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), styles[style])

    flow = [
        paragraph("MEETING PROTOCOL", "Title"),
        paragraph(meeting.title),
        paragraph("Date: " + str(meeting.meeting_date or "Not supplied")),
        paragraph("Participants: " + ", ".join(p.name for p in meeting.participants)),
    ]
    for heading, texts in [
        ("1. Summary", [meeting.summary.text]),
        ("2. Key decisions", [d.text for d in meeting.decisions]),
    ]:
        flow.append(paragraph(heading, "Heading1"))
        flow.extend(paragraph(t) for t in texts)
    flow.append(paragraph("3. Tasks", "Heading1"))
    rows = [
        [paragraph(t) for t in row]
        for row in [["#", "Task", "Assignee", "Deadline", "Status"], *task_rows(meeting)]
    ]
    table = LongTable(rows, colWidths=[23, 195, 90, 95, 70], repeatRows=1, splitInRow=1)
    table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
            ]
        )
    )
    flow += [table, Spacer(1, 12), paragraph("4. Transcript", "Heading1")]
    for segment in meeting.transcript:
        flow.append(
            paragraph(
                f"[{stamp(segment.start)}] {segment.speaker} / {display_name(meeting, segment.speaker)}\n{segment.text}"
            )
        )
        flow.append(Spacer(1, 5))
    stream = BytesIO()
    SimpleDocTemplate(stream, leftMargin=36, rightMargin=36).build(flow)
    return stream.getvalue()
