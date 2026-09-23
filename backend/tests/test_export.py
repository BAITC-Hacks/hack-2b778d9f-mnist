from docx import Document

from app.export import export_draft


def test_export_separates_candidates_and_shows_unknown_deadline(tmp_path, monkeypatch):
    meeting = {"meeting_date": "2026-09-23", "summary": "Қауіпсіздік талқыланды.",
               "roster": ["Алия"], "speaker_names": {"s0": "Алия"},
               "transcript": [{"id": "t1", "start": 3.0, "end": 6.0,
                               "text": "Фиксируем аудит.", "speaker_id": "s0",
                               "speaker_uncertain": False}],
               "actions": [
                   {"text": "Провести аудит", "assignee": "Юр. отдел",
                    "deadline_phrase": None, "due_date": None, "needs_review": False,
                    "source_turn_ids": ["t1"], "confirmation_turn_ids": ["t1"],
                    "confirmation_checked": True},
                   {"text": "Обновить сайт", "assignee": "Алия",
                    "deadline_phrase": "до пятницы", "due_date": "2026-09-25",
                    "source_turn_ids": ["t1"], "confirmation_turn_ids": [],
                    "needs_review": True},
               ]}

    def fake_convert(*_args, **_kwargs):
        (tmp_path / "draft.pdf").write_bytes(b"%PDF demo")

    monkeypatch.setattr("app.export.subprocess.run", fake_convert)
    docx, pdf = export_draft(meeting, tmp_path)
    document = Document(docx)
    paragraphs = " ".join(paragraph.text for paragraph in document.paragraphs)
    cells = " ".join(cell.text for row in document.tables[0].rows for cell in row.cells)
    assert "Черновик протокола" in paragraphs
    assert "Қауіпсіздік" in paragraphs
    assert "На уточнение" in paragraphs
    candidate_cells = " ".join(cell.text for row in document.tables[1].rows for cell in row.cells)
    assert "Обновить сайт" in candidate_cells
    assert "Алия" in candidate_cells
    assert "до пятницы" in candidate_cells
    assert "2026-09-25" in candidate_cells
    assert "Обновить сайт" not in cells
    assert "Юр. отдел" in cells
    assert "не указан — уточнить" in cells
    assert "t1 [00:03.00–00:06.00]: Фиксируем аудит." in cells
    assert "Алия (s0)" in paragraphs
    assert "Транскрипт и источники" in paragraphs
    assert pdf.read_bytes().startswith(b"%PDF")


def test_failed_conversion_preserves_docx_but_not_old_pdf(tmp_path, monkeypatch):
    (tmp_path / "draft.pdf").write_bytes(b"%PDF old")

    def fail(*_args, **_kwargs):
        raise FileNotFoundError("libreoffice")

    monkeypatch.setattr("app.export.subprocess.run", fail)
    try:
        export_draft({"meeting_date": "2026-09-23", "summary": "", "actions": []}, tmp_path)
    except RuntimeError as error:
        assert "PDF conversion failed" in str(error)
    else:
        raise AssertionError("Expected PDF conversion error")
    assert (tmp_path / "draft.docx").exists()
    assert not (tmp_path / "draft.pdf").exists()


def test_docx_only_needs_no_converter_and_missing_review_status_is_candidate(tmp_path, monkeypatch):
    def unexpected(*_args, **_kwargs):
        raise AssertionError("DOCX export must not start a converter")

    monkeypatch.setattr("app.export.subprocess.run", unexpected)
    docx, pdf = export_draft({
        "meeting_date": "2026-09-23", "summary": "Қысқаша", "actions": [
            {"text": "Непроверенное действие", "assignee": "Алия", "deadline_phrase": "ертең"},
        ],
    }, tmp_path, include_pdf=False)
    document = Document(docx)
    assert pdf is None
    assert len(document.tables[0].rows) == 1
    candidates = " ".join(cell.text for row in document.tables[1].rows for cell in row.cells)
    assert "Непроверенное действие" in candidates
    assert "Алия" in candidates
    assert "ертең" in candidates


def test_export_fails_closed_for_incomplete_confirmation(tmp_path, monkeypatch):
    meeting = {"meeting_date": "2026-09-23", "summary": "",
               "transcript": [{"id": "t1", "start": 0, "end": 1, "text": "Подтверждаю"}],
               "actions": [
        {"text": "Missing fields"},
        {"text": "Unchecked", "needs_review": False, "confirmation_checked": True,
         "confirmation_turn_ids": []},
        {"text": "Confirmed", "needs_review": False, "confirmation_checked": True,
         "source_turn_ids": ["t1"], "confirmation_turn_ids": ["t1"],
         "assignee": None, "deadline_phrase": None, "due_date": None},
    ]}
    monkeypatch.setattr("app.export.subprocess.run", lambda *_args, **_kwargs: (tmp_path / "draft.pdf").write_bytes(b"%PDF demo"))

    docx, _ = export_draft(meeting, tmp_path)
    document = Document(docx)
    candidates = " ".join(cell.text for row in document.tables[1].rows for cell in row.cells)
    rows = " ".join(cell.text for row in document.tables[0].rows for cell in row.cells)
    assert "Missing fields" in candidates
    assert "Unchecked" in candidates
    assert "Confirmed" in rows
    assert "Missing fields" not in rows
    assert "Unchecked" not in rows
