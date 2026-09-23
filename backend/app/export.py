"""Render a complete review draft as DOCX, with optional local PDF conversion."""

import subprocess
import tempfile
from datetime import date
from pathlib import Path

from docx import Document


def _time(seconds: float) -> str:
    minutes, remainder = divmod(float(seconds), 60)
    return f"{int(minutes):02d}:{remainder:05.2f}"


def _source(turn_id: str, turns: dict) -> str:
    turn = turns.get(turn_id)
    if turn is None:
        return f"{turn_id} — источник недоступен, уточнить"
    return f"{turn_id} [{_time(turn['start'])}–{_time(turn['end'])}]: {turn['text']}"


def _action_table(document, actions: list[dict], turns: dict) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, title in zip(table.rows[0].cells, ("Действие", "Ответственный", "Срок", "Основания")):
        cell.text = title
    for action in actions:
        cells = table.add_row().cells
        cells[0].text = action.get("text", "")
        cells[1].text = action.get("assignee") or "Не указан — уточнить"
        phrase, due_date = action.get("deadline_phrase"), action.get("due_date")
        cells[2].text = "\n".join((
            f"Как сказано: {phrase}" if phrase else "Срок в записи не указан — уточнить",
            f"Дата исполнения: {due_date}" if due_date else "Дата требует уточнения",
        ))
        sources = [_source(turn_id, turns) for turn_id in action.get("source_turn_ids", [])]
        confirmations = [_source(turn_id, turns) for turn_id in action.get("confirmation_turn_ids", [])]
        cells[3].text = ("Основание:\n" + ("\n".join(sources) or "Не указано — уточнить")
                         + "\nПодтверждение:\n" + ("\n".join(confirmations) or "Не проверено"))


def _confirmed(action, turns: dict) -> bool:
    if not isinstance(action, dict):
        return False
    if (action.get("needs_review") is not False or action.get("confirmation_checked") is not True
            or not isinstance(action.get("text"), str) or not action["text"].strip()):
        return False
    sources, confirmations = action.get("source_turn_ids"), action.get("confirmation_turn_ids")
    if (not isinstance(sources, list) or not sources or any(not isinstance(item, str) for item in sources)
            or not isinstance(confirmations, list) or not confirmations
            or any(not isinstance(item, str) for item in confirmations)
            or not set(sources + confirmations) <= turns.keys()):
        return False
    if any(key not in action or (action[key] is not None and not isinstance(action[key], str))
           for key in ("assignee", "deadline_phrase", "due_date")):
        return False
    if action["due_date"] is not None:
        try:
            if date.fromisoformat(action["due_date"]).isoformat() != action["due_date"]:
                return False
        except ValueError:
            return False
    return True


def export_draft(meeting: dict, output_dir: Path, *, include_pdf: bool = True) -> tuple[Path, Path | None]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    docx_path, pdf_path = output_dir / "draft.docx", output_dir / "draft.pdf"
    pdf_path.unlink(missing_ok=True)
    document = Document()
    document.core_properties.title = "Протокол совещания — Итоги встречи: решения и поручения"
    document.core_properties.author = "Jinalys AI"
    document.add_paragraph("Jinalys AI · От встречи к решениям", style="Subtitle")
    document.add_heading("Черновик протокола", level=0)
    document.add_paragraph("Итоги встречи: решения и поручения", style="Subtitle")
    document.add_paragraph("Документ требует проверки и утверждения уполномоченным лицом.")
    document.add_paragraph(f"Дата совещания: {meeting['meeting_date']}")
    document.add_heading("Участники", level=1)
    document.add_paragraph(", ".join(meeting.get("roster", [])) or "Список участников не указан")
    for speaker_id, name in meeting.get("speaker_names", {}).items():
        document.add_paragraph(f"{speaker_id}: {name} (имя установлено при ручной проверке)")
    document.add_heading("Краткое содержание", level=1)
    document.add_paragraph(meeting.get("summary") or "На уточнение")
    document.add_heading("Поручения", level=1)
    turns = {turn["id"]: turn for turn in meeting.get("transcript", [])}
    confirmed, candidates = [], []
    for action in meeting.get("actions", []):
        checked = _confirmed(action, turns)
        (confirmed if checked else candidates).append(action)
    _action_table(document, confirmed, turns)
    if not confirmed:
        document.add_paragraph("Подтверждённых при проверке поручений нет.")
    if candidates:
        document.add_heading("На уточнение — кандидаты в поручения", level=1)
        _action_table(document, candidates, turns)
    document.add_heading("Транскрипт и источники", level=1)
    for turn_id, turn in turns.items():
        speaker_id = turn.get("speaker_id")
        person = meeting.get("speaker_names", {}).get(speaker_id)
        speaker = f"{person} ({speaker_id})" if person else speaker_id or "Говорящий не установлен"
        flags = []
        if turn.get("overlap"):
            flags.append("одновременная речь; возможны пропуски")
        if turn.get("speaker_uncertain", True):
            flags.append("говорящий требует проверки")
        if turn.get("timestamp_uncertain"):
            flags.append("время приблизительное")
        document.add_paragraph(speaker + (" — " + "; ".join(flags) if flags else ""))
        document.add_paragraph(_source(turn_id, turns))
    document.save(docx_path)
    if not docx_path.is_file() or docx_path.stat().st_size == 0:
        raise RuntimeError("DOCX export is empty")
    if not include_pdf:
        return docx_path, None
    try:
        with tempfile.TemporaryDirectory(prefix="meeting-libreoffice-") as profile:
            subprocess.run(["libreoffice", f"-env:UserInstallation={Path(profile).as_uri()}",
                            "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(docx_path)],
                           check=True, timeout=120, capture_output=True)
    except (OSError, subprocess.CalledProcessError, subprocess.TimeoutExpired) as error:
        raise RuntimeError(f"PDF conversion failed: {error}") from error
    if not pdf_path.is_file() or pdf_path.stat().st_size == 0 or not pdf_path.read_bytes().startswith(b"%PDF"):
        raise RuntimeError("PDF conversion did not produce a valid PDF")
    return docx_path, pdf_path
